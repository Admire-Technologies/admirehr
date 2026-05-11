"""
Export utilities for generating reports in various formats (PDF, Excel, CSV).
"""
import csv
from io import BytesIO, StringIO
from typing import Dict, List, Any
from datetime import datetime
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


class ReportExporter:
    """Base class for report export functionality."""
    
    @staticmethod
    def export_to_csv(data: List[Dict[str, Any]], headers: List[str]) -> str:
        """Export data to CSV format."""
        output = StringIO()
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerows(data)
        return output.getvalue()
    
    @staticmethod
    def export_to_excel(
        data: List[Dict[str, Any]],
        headers: List[str],
        sheet_name: str = 'Report'
    ) -> BytesIO:
        """Export data to Excel format."""
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        
        # Style for headers
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        # Write headers
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        # Write data
        for row_num, row_data in enumerate(data, 2):
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = row_data.get(header, '')

        
        # Auto-adjust column widths
        for col_num, header in enumerate(headers, 1):
            column_letter = get_column_letter(col_num)
            max_length = len(header)
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_num, max_col=col_num):
                for cell in row:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
        
        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
    
    @staticmethod
    def export_to_pdf(
        title: str,
        data: List[Dict[str, Any]],
        headers: List[str],
        summary: Dict[str, Any] = None
    ) -> BytesIO:
        """Export data to PDF format."""
        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        elements = []
        styles = getSampleStyleSheet()
        
        # Title
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#366092'),
            spaceAfter=30,
            alignment=TA_CENTER
        )
        elements.append(Paragraph(title, title_style))
        elements.append(Spacer(1, 0.2 * inch))
        
        # Summary section if provided
        if summary:
            summary_style = ParagraphStyle(
                'Summary',
                parent=styles['Normal'],
                fontSize=10,
                spaceAfter=12
            )
            elements.append(Paragraph('<b>Summary</b>', styles['Heading2']))
            for key, value in summary.items():
                elements.append(Paragraph(f'<b>{key}:</b> {value}', summary_style))
            elements.append(Spacer(1, 0.3 * inch))
        
        # Data table
        if data:
            table_data = [headers]
            for row in data:
                table_data.append([str(row.get(h, '')) for h in headers])

            
            table = Table(table_data)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#366092')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.lightgrey])
            ]))
            elements.append(table)
        
        # Footer
        elements.append(Spacer(1, 0.5 * inch))
        footer_style = ParagraphStyle(
            'Footer',
            parent=styles['Normal'],
            fontSize=8,
            textColor=colors.grey,
            alignment=TA_CENTER
        )
        elements.append(Paragraph(
            f'Generated on {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}',
            footer_style
        ))
        
        doc.build(elements)
        buffer.seek(0)
        return buffer
