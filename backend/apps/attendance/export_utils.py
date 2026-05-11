"""
Attendance export utilities for PDF and Excel generation.
"""

import io
from datetime import datetime
from decimal import Decimal
from typing import List, Dict, Any
from django.http import HttpResponse
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER, TA_LEFT
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


class AttendanceExporter:
    """Base class for attendance export functionality."""
    
    def __init__(self, records: List[Any], summary: Dict[str, Any], company: Any):
        self.records = records
        self.summary = summary
        self.company = company
    
    def export(self) -> HttpResponse:
        """Export attendance data. To be implemented by subclasses."""
        raise NotImplementedError


class AttendancePDFExporter(AttendanceExporter):
    """Export attendance data to PDF format."""
    
    def export(self) -> HttpResponse:
        """Generate PDF report."""
        buffer = io.BytesIO()
        
        # Create PDF document
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        
        # Container for PDF elements
        elements = []
        
        # Styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1a202c'),
            spaceAfter=30,
            alignment=TA_CENTER
        )
        
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#2d3748'),
            spaceAfter=12,
            spaceBefore=12
        )
        
        # Title
        title = Paragraph(f"{self.company.name}<br/>Attendance Report", title_style)
        elements.append(title)
        elements.append(Spacer(1, 12))
        
        # Report metadata
        date_range = self.summary.get('date_range', {})
        metadata_data = [
            ['Report Generated:', datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
            ['Date Range:', f"{date_range.get('start', 'N/A')} to {date_range.get('end', 'N/A')}"],
            ['Total Records:', str(self.summary.get('total_records', 0))],
        ]
        
        metadata_table = Table(metadata_data, colWidths=[2*inch, 4*inch])
        metadata_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#4a5568')),
            ('TEXTCOLOR', (1, 0), (1, -1), colors.HexColor('#1a202c')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))
        elements.append(metadata_table)
        elements.append(Spacer(1, 20))
        
        # Summary section
        summary_heading = Paragraph("Summary Statistics", heading_style)
        elements.append(summary_heading)
        
        summary_data = [
            ['Metric', 'Value'],
            ['Present', str(self.summary.get('present', 0))],
            ['Late', str(self.summary.get('late', 0))],
            ['Absent', str(self.summary.get('absent', 0))],
            ['Half Day', str(self.summary.get('half_day', 0))],
            ['Average Working Hours', f"{self.summary.get('average_working_hours', 0):.2f}"],
            ['Total Working Hours', f"{self.summary.get('total_working_hours', 0):.2f}"],
            ['Total Overtime Hours', f"{self.summary.get('total_overtime_hours', 0):.2f}"],
        ]
        
        summary_table = Table(summary_data, colWidths=[3*inch, 2*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4299e1')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 10),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f7fafc')]),
        ]))
        elements.append(summary_table)
        elements.append(Spacer(1, 20))
        
        # Attendance records section
        if self.records:
            records_heading = Paragraph("Attendance Records", heading_style)
            elements.append(records_heading)
            
            # Table headers
            records_data = [
                ['Date', 'Employee', 'Check In', 'Check Out', 'Hours', 'Status']
            ]
            
            # Add records
            for record in self.records:
                check_in = record.check_in.strftime('%H:%M') if record.check_in else '-'
                check_out = record.check_out.strftime('%H:%M') if record.check_out else '-'
                hours = f"{float(record.working_hours):.2f}" if record.working_hours else '-'
                
                records_data.append([
                    record.date.strftime('%Y-%m-%d'),
                    record.employee.full_name[:20],  # Truncate long names
                    check_in,
                    check_out,
                    hours,
                    record.status.title()
                ])
            
            records_table = Table(records_data, colWidths=[1.2*inch, 1.8*inch, 1*inch, 1*inch, 0.8*inch, 1*inch])
            records_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4299e1')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 9),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f7fafc')]),
            ]))
            elements.append(records_table)
        
        # Build PDF
        doc.build(elements)
        
        # Get PDF data
        pdf_data = buffer.getvalue()
        buffer.close()
        
        # Create response
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="attendance_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf"'
        response.write(pdf_data)
        
        return response


class AttendanceExcelExporter(AttendanceExporter):
    """Export attendance data to Excel format."""
    
    def export(self) -> HttpResponse:
        """Generate Excel report."""
        # Create workbook
        wb = openpyxl.Workbook()
        
        # Remove default sheet
        wb.remove(wb.active)
        
        # Create summary sheet
        self._create_summary_sheet(wb)
        
        # Create records sheet
        self._create_records_sheet(wb)
        
        # Save to buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        # Create response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="attendance_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        response.write(buffer.getvalue())
        
        return response
    
    def _create_summary_sheet(self, wb):
        """Create summary statistics sheet."""
        ws = wb.create_sheet('Summary')
        
        # Header style
        header_fill = PatternFill(start_color='4299E1', end_color='4299E1', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=12)
        
        # Title
        ws['A1'] = f'{self.company.name} - Attendance Report'
        ws['A1'].font = Font(bold=True, size=16)
        ws.merge_cells('A1:B1')
        
        # Report metadata
        ws['A3'] = 'Report Generated:'
        ws['B3'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        
        date_range = self.summary.get('date_range', {})
        ws['A4'] = 'Date Range:'
        ws['B4'] = f"{date_range.get('start', 'N/A')} to {date_range.get('end', 'N/A')}"
        
        ws['A5'] = 'Total Records:'
        ws['B5'] = self.summary.get('total_records', 0)
        
        # Summary statistics
        ws['A7'] = 'Metric'
        ws['B7'] = 'Value'
        ws['A7'].fill = header_fill
        ws['B7'].fill = header_fill
        ws['A7'].font = header_font
        ws['B7'].font = header_font
        
        summary_data = [
            ('Present', self.summary.get('present', 0)),
            ('Late', self.summary.get('late', 0)),
            ('Absent', self.summary.get('absent', 0)),
            ('Half Day', self.summary.get('half_day', 0)),
            ('Average Working Hours', f"{self.summary.get('average_working_hours', 0):.2f}"),
            ('Total Working Hours', f"{self.summary.get('total_working_hours', 0):.2f}"),
            ('Total Overtime Hours', f"{self.summary.get('total_overtime_hours', 0):.2f}"),
            ('Unique Employees', self.summary.get('unique_employees', 0)),
        ]
        
        for idx, (metric, value) in enumerate(summary_data, start=8):
            ws[f'A{idx}'] = metric
            ws[f'B{idx}'] = value
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 20
    
    def _create_records_sheet(self, wb):
        """Create attendance records sheet."""
        ws = wb.create_sheet('Attendance Records')
        
        # Header style
        header_fill = PatternFill(start_color='4299E1', end_color='4299E1', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        # Border style
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Headers
        headers = ['Date', 'Employee ID', 'Employee Name', 'Department', 'Check In', 'Check Out', 'Working Hours', 'Overtime', 'Status', 'Biometric']
        for col_num, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Add records
        for row_num, record in enumerate(self.records, start=2):
            check_in = record.check_in.strftime('%H:%M:%S') if record.check_in else '-'
            check_out = record.check_out.strftime('%H:%M:%S') if record.check_out else '-'
            hours = float(record.working_hours) if record.working_hours else 0
            
            # Calculate overtime
            overtime_threshold = record.company.overtime_threshold_hours
            overtime = max(0, hours - float(overtime_threshold))
            
            row_data = [
                record.date.strftime('%Y-%m-%d'),
                record.employee.employee_id,
                record.employee.full_name,
                record.employee.department.name if record.employee.department else '-',
                check_in,
                check_out,
                hours,
                overtime,
                record.status.title(),
                'Yes' if record.biometric_verified else 'No'
            ]
            
            for col_num, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = value
                cell.border = thin_border
                
                # Center alignment for specific columns
                if col_num in [1, 5, 6, 9, 10]:
                    cell.alignment = Alignment(horizontal='center')
        
        # Adjust column widths
        column_widths = [12, 15, 25, 20, 12, 12, 15, 12, 12, 12]
        for col_num, width in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(col_num)].width = width
        
        # Freeze header row
        ws.freeze_panes = 'A2'


def export_attendance_report(records: List[Any], summary: Dict[str, Any], company: Any, format: str) -> HttpResponse:
    """
    Export attendance report in specified format.
    
    Args:
        records: List of AttendanceRecord instances
        summary: Dictionary with summary statistics
        company: Company instance
        format: Export format ('pdf' or 'excel')
    
    Returns:
        HttpResponse with exported file
    """
    if format == 'pdf':
        exporter = AttendancePDFExporter(records, summary, company)
    elif format == 'excel':
        exporter = AttendanceExcelExporter(records, summary, company)
    else:
        raise ValueError(f"Unsupported export format: {format}")
    
    return exporter.export()
